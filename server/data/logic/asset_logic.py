"""
资产CRUD、状态流转、乐观锁、变更日志 业务逻辑
不得出现任何SQL语句
"""
from typing import Any, Dict, List, Optional, Tuple
from datetime import datetime
import openpyxl
from pathlib import Path

from data.basic.asset_dao import AssetDAO, CategoryDAO, AssetOperLogDAO


class AssetLogic:
    """资产业务逻辑"""

    def __init__(self, db_config: Dict):
        self.asset_dao = AssetDAO(db_config)
        self.category_dao = CategoryDAO(db_config)
        self.oper_log_dao = AssetOperLogDAO(db_config)

    def list_assets(self, keyword: str = "", category_id: int = 0,
                    status: int = 0, dept_id: int = 0,
                    page_num: int = 1, page_size: int = 10) -> Tuple[List, int]:
        return self.asset_dao.list_assets(keyword, category_id, status, dept_id, page_num, page_size)

    def get_asset_detail(self, asset_id: int) -> Optional[Dict]:
        asset = self.asset_dao.get_asset_by_id(asset_id)
        if not asset:
            return None
        logs = self.oper_log_dao.list_logs_by_asset(asset_id)
        asset["logs"] = logs
        return asset

    def create_asset(self, data: Dict, operator_id: int) -> Dict:
        if not data.get("asset_code"):
            data["asset_code"] = self.asset_dao.generate_asset_code()
        # DAO 必填字段兜底，避免 SQLAlchemy 报 InvalidRequestError
        data.setdefault("status", 1)
        data.setdefault("spec", "")
        data.setdefault("category_id", 0)
        data.setdefault("purchase_date", None)
        data.setdefault("book_value", 0)
        data.setdefault("location", "")
        data.setdefault("dept_id", 0)
        data.setdefault("user_id", 0)
        data.setdefault("remark", "")
        asset_id = self.asset_dao.insert_asset(data)
        # 记录操作日志
        self.oper_log_dao.insert_log({
            "asset_id": asset_id,
            "operator_id": operator_id,
            "oper_type": "CREATE",
            "field_name": "ALL",
            "old_value": "",
            "new_value": f"创建资产 {data.get('asset_name')}",
            "oper_ip": "127.0.0.1",
        })
        return self.get_asset_detail(asset_id)

    def update_asset(self, data: Dict, operator_id: int, oper_ip: str = "127.0.0.1") -> Optional[Dict]:
        """更新资产，乐观锁校验，记录变更日志"""
        old = self.asset_dao.get_asset_by_id(data["asset_id"])
        if not old:
            return None
        # 乐观锁校验
        if old["version"] != data.get("version", old["version"]):
            return {"error": "version_conflict", "current_version": old["version"]}

        # DAO 必填字段兜底：以前端提交为准，未提交则保留原值
        default_fields = [
            "asset_name", "spec", "category_id", "purchase_date",
            "book_value", "location", "dept_id", "user_id", "status", "remark",
        ]
        for f in default_fields:
            data.setdefault(f, old.get(f, "" if f in ("asset_name", "spec", "location", "remark") else 0))
        # 关键：version 是 DAO update_asset SQL 必填绑定参数，前端 form 未携带
        data.setdefault("version", old["version"])

        # 比较变更字段
        compare_fields = [
            ("asset_name", "资产名称"), ("spec", "规格型号"),
            ("category_id", "资产分类"), ("purchase_date", "购置日期"),
            ("book_value", "账面价值"), ("location", "存放地点"),
            ("dept_id", "归属学院"), ("user_id", "使用人"),
            ("status", "资产状态"), ("remark", "备注"),
        ]
        for field, field_cn in compare_fields:
            old_val = str(old.get(field, ""))
            new_val = str(data.get(field, old.get(field, "")))
            if old_val != new_val:
                self.oper_log_dao.insert_log({
                    "asset_id": data["asset_id"],
                    "operator_id": operator_id,
                    "oper_type": "UPDATE",
                    "field_name": field_cn,
                    "old_value": old_val,
                    "new_value": new_val,
                    "oper_ip": oper_ip,
                })

        self.asset_dao.update_asset(data)
        return self.get_asset_detail(data["asset_id"])

    def update_asset_status(self, asset_id: int, status_id: int,
                            user_id: int = 0, version: int = 0) -> bool:
        result = self.asset_dao.update_asset_status(asset_id, status_id, user_id, version)
        return result > 0

    def delete_asset(self, asset_id: int) -> bool:
        self.asset_dao.delete_asset(asset_id)
        return True

    def get_category_list(self) -> List[Dict]:
        categories = self.category_dao.list_categories()
        return self._build_category_tree(categories, 0)

    def _build_category_tree(self, categories: List[Dict], parent_id: int) -> List[Dict]:
        tree = []
        for c in categories:
            if c.get("parent_id", 0) == parent_id:
                node = {**c, "children": self._build_category_tree(categories, c["category_id"])}
                tree.append(node)
        return tree

    def get_status_list(self) -> List[Dict]:
        return self.asset_dao.list_all_asset_status()

    def export_assets(self, keyword: str = "", category_id: int = 0,
                      status: int = 0, dept_id: int = 0,
                      output_path: str = "") -> str:
        """导出资产为Excel"""
        data = self.asset_dao.get_export_list(keyword, category_id, status, dept_id)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "资产台账"
        headers = ["资产编号", "资产名称", "规格型号", "分类", "购置日期", "账面价值",
                    "存放地点", "归属学院", "使用人", "状态", "备注"]
        ws.append(headers)
        for row in data:
            ws.append([
                row.get("asset_code", ""), row.get("asset_name", ""),
                row.get("spec", ""), row.get("category_name", ""),
                str(row.get("purchase_date", "")), row.get("book_value", 0),
                row.get("location", ""), row.get("dept_name", ""),
                row.get("user_name", ""), row.get("status_name", ""),
                row.get("remark", ""),
            ])
        if not output_path:
            output_path = str(Path.cwd() / "temp" / f"asset_export_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")
        wb.save(output_path)
        return output_path

    def import_assets(self, file_path: str, operator_id: int) -> Dict:
        """批量导入资产"""
        status_list = self.get_status_list()
        status_name_map = {s["status_name"]: s["status_id"] for s in status_list}

        wb = openpyxl.load_workbook(file_path)
        ws = wb.active
        success = 0
        fail = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            try:
                if not row[1]:
                    continue
                status_text = str(row[9]).strip() if len(row) > 9 and row[9] else "闲置中"
                asset_status = status_name_map.get(status_text, 1)
                data = {
                    "asset_code": self.asset_dao.generate_asset_code(),
                    "asset_name": str(row[1]) if row[1] else "",
                    "spec": str(row[2]) if len(row) > 2 and row[2] else "",
                    "category_id": int(row[3]) if len(row) > 3 and row[3] else 0,
                    "purchase_date": str(row[4]) if len(row) > 4 and row[4] else None,
                    "book_value": float(row[5]) if len(row) > 5 and row[5] else 0,
                    "location": str(row[6]) if len(row) > 6 and row[6] else "",
                    "dept_id": int(row[7]) if len(row) > 7 and row[7] else 0,
                    "user_id": int(row[8]) if len(row) > 8 and row[8] else 0,
                    "status": asset_status,
                    "remark": str(row[10]) if len(row) > 10 and row[10] else "",
                }
                self.create_asset(data, operator_id)
                success += 1
            except Exception:
                fail += 1
        return {"success_count": success, "fail_count": fail}

    def get_oper_logs(self, page_num: int = 1, page_size: int = 20) -> tuple:
        # 返回分页的操作日志
        return [], 0  # 简化实现