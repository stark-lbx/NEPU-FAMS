"""
统计报表 业务逻辑 — 跨域聚合，在 Logic 层完成跨库数据关联
不得出现任何SQL语句
"""
from typing import Dict, List
import tempfile
import os
import datetime

from data.basic.report_dao import ReportAssetDAO, ReportRepairDAO, ReportInventoryDAO, ReportApplyDAO
from data.basic.user_dao import DeptDAO


class ReportLogic:
    """跨域统计报表业务逻辑"""

    def __init__(self, user_db_config: Dict, asset_db_config: Dict,
                 repair_db_config: Dict, inventory_db_config: Dict,
                 workflow_db_config: Dict):
        self.dept_dao = DeptDAO(user_db_config)
        self.asset_dao = ReportAssetDAO(asset_db_config)
        self.repair_dao = ReportRepairDAO(repair_db_config)
        self.inventory_dao = ReportInventoryDAO(inventory_db_config)
        self.apply_dao = ReportApplyDAO(workflow_db_config)

    # ---- 内部联表方法 ----

    def _resolve_dept_names(self, rows: List[Dict]) -> List[Dict]:
        """将 dept_id 列表转换为含 dept_name 的结果"""
        if not rows:
            return []
        dept_ids = list({r["dept_id"] for r in rows if r.get("dept_id")})
        dept_map = self.dept_dao.get_dept_map(dept_ids) if dept_ids else {}
        return [
            {"label": dept_map.get(r["dept_id"], f"部门#{r['dept_id']}"), "value": r["count"]}
            for r in rows
        ]

    # ---- 资产报表 ----

    def get_asset_report(self, dept_id: int = 0) -> Dict:
        total = self.asset_dao.total_count(dept_id)
        dept_raw = self.asset_dao.dept_asset_count(dept_id)
        dept_chart = self._resolve_dept_names(dept_raw)
        return {
            "total": total,
            "dept_chart": dept_chart,
            "status_pie": self.asset_dao.status_distribution(dept_id),
            "category_pie": self.asset_dao.category_distribution(dept_id),
            "monthly_new": self.asset_dao.monthly_new_assets(dept_id),
        }

    # ---- 维修报表 ----

    def get_repair_report(self, dept_id: int = 0) -> Dict:
        total = self.repair_dao.total_count(dept_id)
        dept_raw = self.repair_dao.dept_repair_count(dept_id)
        dept_chart = self._resolve_dept_names(dept_raw)
        return {
            "total": total,
            "monthly": self.repair_dao.monthly_stats(dept_id),
            "status_summary": self.repair_dao.status_summary(dept_id),
            "dept_chart": dept_chart,
        }

    # ---- 盘点报表 ----

    def get_inventory_report(self, dept_id: int = 0) -> Dict:
        total = self.inventory_dao.total_task_count(dept_id)
        return {
            "total": total,
            "quarterly": self.inventory_dao.quarterly_diff(dept_id),
            "status_summary": self.inventory_dao.task_completion_stats(dept_id),
        }

    # ---- 仪表盘概览 ----

    def get_dashboard_overview(self, dept_id: int = 0) -> Dict:
        return {
            "asset_total": self.asset_dao.total_count(dept_id),
            "pending_repairs": self.repair_dao.pending_count(dept_id),
            "in_progress_repairs": self.repair_dao.in_progress_count(dept_id),
            "pending_acceptance": self.repair_dao.pending_acceptance_count(dept_id),
            "completed_this_month": self.repair_dao.completed_this_month_count(dept_id),
            "inventory_total": self.inventory_dao.total_task_count(dept_id),
            "pending_applies": self.apply_dao.pending_count(dept_id),
        }

    # ---- 导出 ----

    def export_report(self) -> str:
        try:
            import openpyxl
        except ImportError:
            return ""

        wb = openpyxl.Workbook()

        # Sheet 1: 资产统计
        ws1 = wb.active
        ws1.title = "资产统计"
        ws1.append(["学院", "资产数量"])
        for row in self.get_asset_report()["dept_chart"]:
            ws1.append([row["label"], row["value"]])

        # Sheet 2: 维修统计
        ws2 = wb.create_sheet("维修统计")
        ws2.append(["月份", "提交数", "完成数"])
        for row in self.repair_dao.monthly_stats():
            ws2.append([row.get("month", ""), row.get("submitted", 0), row.get("completed", 0)])

        # Sheet 3: 盘点统计
        ws3 = wb.create_sheet("盘点统计")
        ws3.append(["任务", "盘盈", "盘亏"])
        for row in self.inventory_dao.quarterly_diff():
            ws3.append([row.get("label", ""), row.get("surplus", 0), row.get("shortage", 0)])

        output_path = os.path.join(
            tempfile.gettempdir(),
            f"FAMS_report_{datetime.datetime.now().strftime('%Y%m%d%H%M%S')}.xlsx"
        )
        wb.save(output_path)
        return output_path
